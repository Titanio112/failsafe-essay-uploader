"""Leitura do XLSX, listagem de JPGs, normalizacao e conciliacao.

Estados de cada arquivo conciliado:

- ``EXATO_UNICO``: correspondencia normalizada exata e unica com um aluno.
- ``SUGESTAO_FUZZY``: melhor candidato aproximado acima do limiar e com
  margem para o segundo. NUNCA entra na fila real sem revisao humana.
- ``AMBIGUO``: correspondencia exata com mais de um aluno, ou fuzzy com
  margem insuficiente / mais de um candidato no limiar.
- ``ARQUIVO_SEM_ALUNO``: nenhum candidato acima do limiar.
- ``ALUNO_SEM_ARQUIVO``: aluno da sala sem nenhum arquivo associado.
- ``SEM_CODIGO``: arquivo sem código no nome nem na pasta.
- ``CODIGO_INVALIDO``: código extraído não existe no combo do site.
- ``REVISAO_MANUAL``: casos que exigem revisão humana (duplicatas, conflitos, etc.)
"""

from __future__ import annotations

import re
import unicodedata
import warnings
from dataclasses import dataclass
from pathlib import Path
from typing import Optional, Set

import openpyxl
from rapidfuzz import fuzz

from config import Config

EXATO_UNICO = "EXATO_UNICO"
SUGESTAO_FUZZY = "SUGESTAO_FUZZY"
AMBIGUO = "AMBIGUO"
ARQUIVO_SEM_ALUNO = "ARQUIVO_SEM_ALUNO"
ALUNO_SEM_ARQUIVO = "ALUNO_SEM_ARQUIVO"
SEM_CODIGO = "SEM_CODIGO"
CODIGO_INVALIDO = "CODIGO_INVALIDO"
REVISAO_MANUAL = "REVISAO_MANUAL"

ESTADOS_CONCILIACAO = (
    EXATO_UNICO,
    SUGESTAO_FUZZY,
    AMBIGUO,
    ARQUIVO_SEM_ALUNO,
    ALUNO_SEM_ARQUIVO,
    SEM_CODIGO,
    CODIGO_INVALIDO,
    REVISAO_MANUAL,
)

MARCADOR_PESSOA = " - Red"

class PlanilhaError(RuntimeError):
    """Planilha ausente ou sem as colunas esperadas."""

_PONTUACAO_RE = re.compile(r"[^\w\s]", re.UNICODE)
_ESPACOS_RE = re.compile(r"\s+")

PADRAO_CODIGO_ARQUIVO = re.compile(r"^\s*(\d+)(?=[\s_.-]|$)")
PADRAO_CODIGO_TEMA = re.compile(r"^\s*(\d+)(?=\s*-|\s|$)")

def normalizar_nome(texto: str) -> str:
    if not texto:
        return ""
    decomposto = unicodedata.normalize("NFKD", str(texto))
    sem_acento = "".join(c for c in decomposto if not unicodedata.combining(c))
    minusculo = sem_acento.casefold()
    sem_pontuacao = _PONTUACAO_RE.sub(" ", minusculo)
    sem_apostrofo = sem_pontuacao.replace("'", " ").replace("'", " ")
    return _ESPACOS_RE.sub(" ", sem_apostrofo).strip()

def extrair_pessoa(nome_arquivo: str) -> str:
    stem = Path(nome_arquivo).stem
    codigo_match = PADRAO_CODIGO_ARQUIVO.match(stem)
    if codigo_match:
        pos_fim_codigo = codigo_match.end()
        resto = stem[pos_fim_codigo:].lstrip(" _.-")
    else:
        resto = stem
    pessoa, sep, _ = resto.partition(MARCADOR_PESSOA)
    return pessoa.strip() if sep else resto.strip()

def extrair_codigo_tema(nome_arquivo: str, pasta: str = "") -> Optional[str]:
    """Extrai código do tema do arquivo. Se houver conflito arquivo vs pasta, retorna None para sinalizar REVISAO_MANUAL."""
    stem = Path(nome_arquivo).stem
    match_arquivo = PADRAO_CODIGO_ARQUIVO.match(stem)
    codigo_arquivo = match_arquivo.group(1) if match_arquivo else None
    
    codigo_pasta = None
    if pasta:
        match_pasta = PADRAO_CODIGO_ARQUIVO.match(pasta)
        if match_pasta:
            codigo_pasta = match_pasta.group(1)
    
    # Conflito: ambos existem e são diferentes
    if codigo_arquivo and codigo_pasta and codigo_arquivo != codigo_pasta:
        return None  # Sinaliza conflito -> REVISAO_MANUAL no conciliar
    
    # Prioridade: arquivo > pasta
    return codigo_arquivo or codigo_pasta

def extrair_codigo_do_tema_site(texto_tema: str) -> Optional[str]:
    match = PADRAO_CODIGO_TEMA.match(texto_tema)
    if match:
        return match.group(1)
    return None

def validar_codigo_site(codigo: str, codigos_site: Set[str]) -> bool:
    return codigo in codigos_site

def mascarar_email(email: str) -> str:
    if not email or "@" not in email:
        return "***"
    local, dominio = email.split("@", 1)
    return f"{local[:2]}***@{dominio}" if local else f"***@{dominio}"

@dataclass
class Aluno:
    nome: str
    email: str
    sala: str
    creditos: int | None = None

    @property
    def nome_norm(self) -> str:
        return normalizar_nome(self.nome)

    @property
    def id(self) -> str:
        return self.email.casefold() if self.email else self.nome_norm

@dataclass
class ItemConciliacao:
    pasta: str
    arquivo: str
    pessoa: str
    estado: str
    aluno: Aluno | None = None
    candidatos: list | None = None
    score_fuzzy: float = 0.0
    segundo_score: float = 0.0
    observacao: str = ""
    codigo_tema: str = ""
    codigo_valido: bool = True

    @property
    def chave(self) -> str:
        return f"{self.pasta}/{self.arquivo}"

@dataclass
class ResultadoConciliacao:
    itens: list

    def por_estado(self, estado: str) -> list:
        return [i for i in self.itens if i.estado == estado]

    def fila_validada(self) -> list:
        return [i for i in self.itens if i.estado == EXATO_UNICO]

    def resumo(self) -> dict:
        return {e: len(self.por_estado(e)) for e in ESTADOS_CONCILIACAO}

def _sugerir_fuzzy(norm: str, alunos: list, limiar: float, margem: float):
    pontuacoes = sorted(
        ((fuzz.token_set_ratio(norm, a.nome_norm), a) for a in alunos),
        key=lambda t: t[0],
        reverse=True,
    )
    if not pontuacoes or pontuacoes[0][0] < limiar:
        return ARQUIVO_SEM_ALUNO, None, [], 0.0, 0.0
    melhor_score, melhor_aluno = pontuacoes[0]
    segundo_score = pontuacoes[1][0] if len(pontuacoes) > 1 else 0.0
    if melhor_score - segundo_score < margem:
        candidatos = [a for s, a in pontuacoes if s >= limiar]
        return AMBIGUO, None, candidatos, melhor_score, segundo_score
    return SUGESTAO_FUZZY, melhor_aluno, [melhor_aluno], melhor_score, segundo_score

def _emitir_aviso(mensagem: str) -> None:
    aviso = f"[AVISO] {mensagem}"
    print(aviso)
    warnings.warn(aviso, UserWarning, stacklevel=3)

def carregar_alunos(cfg: Config) -> list[Aluno]:
    wb = openpyxl.load_workbook(cfg.planilha, data_only=True)
    ws = wb[cfg.aba_planilha] if cfg.aba_planilha else wb.active

    header = [str(c.value).strip().lower() if c.value else "" for c in next(ws.iter_rows(min_row=1, max_row=1))]
    try:
        idx_nome = header.index(cfg.coluna_nome.lower())
        idx_email = header.index(cfg.coluna_email.lower())
        idx_sala = header.index(cfg.coluna_sala.lower())
    except ValueError as e:
        raise PlanilhaError(f"Coluna não encontrada: {e}")

    alunos = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if len(row) <= max(idx_nome, idx_email, idx_sala):
            continue
        nome = str(row[idx_nome]).strip() if row[idx_nome] else ""
        email = str(row[idx_email]).strip() if row[idx_email] else ""
        sala = str(row[idx_sala]).strip() if row[idx_sala] else ""
        if not nome or not email or not sala:
            continue
        alunos.append(Aluno(nome=nome, email=email, sala=sala))
    return alunos


def descobrir_pastas(cfg: Config) -> list[Path]:
    raiz = cfg.pasta_raiz
    if not raiz.exists():
        return []
    return sorted([p for p in raiz.iterdir() if p.is_dir() and not p.name.startswith(".")])


def descobrir_arquivos(cfg: Config) -> dict[str, list[str]]:
    extensoes_permitidas = {ext.lower() for ext in cfg.extensoes}
    resultado = {}
    for pasta in descobrir_pastas(cfg):
        arquivos = []
        for f in sorted(pasta.iterdir(), key=lambda x: x.name):
            if not f.is_file():
                continue
            ext = f.suffix.lower()
            if ext not in extensoes_permitidas:
                _emitir_aviso(f"Arquivo ignorado (extensão não permitida): {pasta.name}/{f.name} (ext: {ext})")
                continue
            arquivos.append(f.name)
        if arquivos:
            resultado[pasta.name] = arquivos
    return resultado


def conciliar(
    alunos: list[Aluno],
    arquivos_por_pasta: dict[str, list[str]],
    codigos_site: Set[str],
    limiar: float = 95.0,
    margem: float = 8.0,
) -> ResultadoConciliacao:
    por_norm = {}
    for a in alunos:
        por_norm.setdefault(a.nome_norm, []).append(a)

    # Contar quantos arquivos cada aluno tem por código (antes de processar)
    # para detectar duplicatas de mesmo aluno + mesmo código
    from collections import defaultdict
    contagem_aluno_codigo: dict[tuple[str, str], int] = defaultdict(int)
    
    # Primeira passada: contar
    for pasta, arquivos in arquivos_por_pasta.items():
        for nome_arquivo in arquivos:
            pessoa = extrair_pessoa(nome_arquivo)
            norm = normalizar_nome(pessoa)
            codigo = extrair_codigo_tema(nome_arquivo, pasta)
            if codigo and codigo != "":
                exatos = por_norm.get(norm, [])
                if len(exatos) == 1:
                    aluno_match = exatos[0]
                    key = (aluno_match.id, codigo)
                    contagem_aluno_codigo[key] += 1

    itens = []
    vistos_por_aluno_codigo: dict[tuple[str, str], int] = {}

    for pasta, arquivos in arquivos_por_pasta.items():
        for nome_arquivo in arquivos:
            pessoa = extrair_pessoa(nome_arquivo)
            norm = normalizar_nome(pessoa)

            # Extrair códigos do arquivo e da pasta separadamente para detectar conflito
            stem = Path(nome_arquivo).stem
            match_arquivo = PADRAO_CODIGO_ARQUIVO.match(stem)
            codigo_arquivo = match_arquivo.group(1) if match_arquivo else None
            
            codigo_pasta = None
            if pasta:
                match_pasta = PADRAO_CODIGO_ARQUIVO.match(pasta)
                if match_pasta:
                    codigo_pasta = match_pasta.group(1)
            
            # Detectar conflito: ambos existem e são diferentes
            conflito_pasta_arquivo = (
                codigo_arquivo is not None and 
                codigo_pasta is not None and 
                codigo_arquivo != codigo_pasta
            )
            
            # Código efetivo (prioridade: arquivo > pasta)
            codigo = codigo_arquivo or codigo_pasta
            codigo_valido = True
            estado_inicial = ""

            if conflito_pasta_arquivo:
                estado_inicial = REVISAO_MANUAL
                codigo_valido = False
            elif codigo is None:
                estado_inicial = SEM_CODIGO
                codigo_valido = False
            elif codigos_site and codigo not in codigos_site:
                estado_inicial = CODIGO_INVALIDO
                codigo_valido = False
            else:
                estado_inicial = ""

            exatos = por_norm.get(norm, [])

            if len(exatos) == 1:
                aluno_match = exatos[0]
                # Verificar se há duplicatas para este aluno+código
                key = (aluno_match.id, codigo) if codigo and codigo != "" else None
                tem_duplicata = key and contagem_aluno_codigo.get(key, 0) > 1
                
                if tem_duplicata:
                    estado = REVISAO_MANUAL
                else:
                    estado = EXATO_UNICO
                    if key:
                        vistos_por_aluno_codigo[key] = 1
                
                # Sobrescrever com estado_inicial se houver (SEM_CODIGO, CODIGO_INVALIDO)
                if estado_inicial:
                    estado = estado_inicial

                itens.append(ItemConciliacao(
                    pasta=pasta, arquivo=nome_arquivo, pessoa=pessoa,
                    estado=estado, aluno=aluno_match,
                    codigo_tema=codigo or "", codigo_valido=codigo_valido,
                ))

            elif len(exatos) > 1:
                itens.append(ItemConciliacao(
                    pasta=pasta, arquivo=nome_arquivo, pessoa=pessoa,
                    estado=AMBIGUO, candidatos=list(exatos),
                    observacao=f"{len(exatos)} alunos com o mesmo nome normalizado",
                    codigo_tema=codigo or "", codigo_valido=codigo_valido,
                ))

            else:
                estado, sugerido, candidatos, melhor, segundo = _sugerir_fuzzy(
                    norm, alunos, limiar, margem
                )
                obs = ""
                if estado == SUGESTAO_FUZZY:
                    obs = "revisao humana obrigatoria antes de validar"
                    if sugerido is not None:
                        if codigo and codigo != "":
                            key = (sugerido.id, codigo)
                            if key in vistos_por_aluno_codigo:
                                estado = REVISAO_MANUAL
                            else:
                                vistos_por_aluno_codigo[key] = 1
                elif estado == AMBIGUO:
                    obs = "candidatos fuzzy sem margem suficiente"
                if estado_inicial:
                    estado = estado_inicial
                    obs = obs or "codigo ausente ou invalido"

                itens.append(ItemConciliacao(
                    pasta=pasta, arquivo=nome_arquivo, pessoa=pessoa,
                    estado=estado, aluno=sugerido, candidatos=candidatos,
                    score_fuzzy=round(melhor, 2), segundo_score=round(segundo, 2),
                    observacao=obs,
                    codigo_tema=codigo or "", codigo_valido=codigo_valido,
                ))

    alunos_com_arquivo = set()
    for item in itens:
        if item.aluno:
            alunos_com_arquivo.add(id(item.aluno))

    for a in alunos:
        if id(a) not in alunos_com_arquivo:
            itens.append(ItemConciliacao(
                pasta="", arquivo="", pessoa=a.nome,
                estado=ALUNO_SEM_ARQUIVO, aluno=a,
            ))

    return ResultadoConciliacao(itens=itens)